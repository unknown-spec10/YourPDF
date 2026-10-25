"""Synchronous office processing functions (in-request, no queue)."""

import logging
import os
import shutil
import subprocess
import sys
import uuid
import zipfile

from pdf2image import convert_from_path

from app.tasks.common import (
    SUBPROCESS_TIMEOUT,
    SUBPROCESS_TIMEOUT_MESSAGE,
    cleanup_paths,
    ensure_tmp_dir,
    read_bytes,
    resolve_path,
    resolve_paths,
)

logger = logging.getLogger(__name__)


def _download_name(original_filename: str | None, fallback: str, build) -> str:
    if original_filename:
        base, ext = os.path.splitext(original_filename)
        return build(base, ext)
    return fallback


def _convert_docx_or_pptx_to_pdf(file_path: str, output_dir: str) -> str:
    """
    Converts a document (DOCX/PPTX) to PDF using LibreOffice headless mode.
    Returns the path to the generated PDF.
    """
    soffice_cmd = None
    if sys.platform == "win32":
        paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        for p in paths:
            if os.path.exists(p):
                soffice_cmd = p
                break

    if not soffice_cmd:
        for cmd in ["soffice", "libreoffice"]:
            if shutil.which(cmd):
                soffice_cmd = cmd
                break

    if not soffice_cmd:
        raise Exception("LibreOffice soffice command not found. Cannot perform document conversion.")

    cmd = [soffice_cmd, "--headless", "--convert-to", "pdf", "--outdir", output_dir, file_path]
    logger.info(f"Running LibreOffice conversion command: {' '.join(cmd)}")
    try:
        subprocess.run(cmd, capture_output=True, text=True, check=True, timeout=SUBPROCESS_TIMEOUT)
    except subprocess.TimeoutExpired:
        logger.error("LibreOffice conversion timed out after %s seconds", SUBPROCESS_TIMEOUT)
        raise Exception(SUBPROCESS_TIMEOUT_MESSAGE)

    base_name = os.path.splitext(os.path.basename(file_path))[0]
    generated_pdf = os.path.join(output_dir, f"{base_name}.pdf")
    if not os.path.exists(generated_pdf):
        raise FileNotFoundError(f"LibreOffice execution completed, but output PDF was not found at {generated_pdf}")

    return generated_pdf


def merge_docx_task(file_paths: list[str], original_filename: str = None) -> tuple[bytes, str]:
    """Combine Word documents into one file and return the bytes."""
    logger.info(f"Starting DOCX merge task for files: {file_paths}")

    from docx import Document

    resolved_paths = resolve_paths(file_paths)

    output_filename = f"merged_{uuid.uuid4()}.docx"
    tmp_dir = ensure_tmp_dir()
    output_path = os.path.join(tmp_dir, output_filename)

    try:
        valid_paths = [fp for fp in resolved_paths if os.path.exists(fp)]
        if not valid_paths:
            raise Exception("No valid Word documents found to merge.")

        merged_document = Document(valid_paths[0])

        for fp in valid_paths[1:]:
            merged_document.add_page_break()
            sub_doc = Document(fp)

            for element in sub_doc.element.body:
                if element.tag.endswith("sectPr"):
                    continue
                merged_document.element.body.append(element)

        merged_document.save(output_path)

        data = read_bytes(output_path)
        download_name = _download_name(
            original_filename, output_filename, lambda base, ext: f"{base}_merged{ext}"
        )
        return data, download_name
    except Exception as e:
        logger.error(f"DOCX merge task failed: {e}")
        raise e
    finally:
        for fp in resolved_paths:
            if fp and os.path.exists(fp):
                try:
                    os.remove(fp)
                except Exception:
                    pass
        if os.path.exists(output_path):
            try:
                os.remove(output_path)
            except Exception:
                pass


def docx_to_images_task(
    file_path: str, img_format: str, dpi: int, original_filename: str = None
) -> tuple[bytes, str]:
    """Convert DOCX to PDF first, rasterize pages, return single image or ZIP bytes."""
    logger.info(f"Starting DOCX to image task. Format: {img_format}, DPI: {dpi}")

    file_path = resolve_path(file_path)
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Input file not found: {file_path}")

    tmp_dir = ensure_tmp_dir()

    pdf_path = None
    output_filename = ""
    output_path = ""

    try:
        pdf_path = _convert_docx_or_pptx_to_pdf(file_path, tmp_dir)

        images = convert_from_path(pdf_path, dpi=dpi)

        if len(images) == 1:
            output_filename = f"image_{uuid.uuid4()}.{img_format.lower()}"
            output_path = os.path.join(tmp_dir, output_filename)
            images[0].save(output_path, format=img_format.upper())

            download_name = _download_name(
                original_filename,
                output_filename,
                lambda base, ext: f"{base}.{img_format.lower()}",
            )
        else:
            output_filename = f"images_{uuid.uuid4()}.zip"
            output_path = os.path.join(tmp_dir, output_filename)
            with zipfile.ZipFile(output_path, "w") as zipf:
                for idx, img in enumerate(images):
                    img_name = f"page_{idx + 1}.{img_format.lower()}"
                    temp_img_path = os.path.join(tmp_dir, img_name)
                    img.save(temp_img_path, format=img_format.upper())
                    zipf.write(temp_img_path, img_name)
                    os.remove(temp_img_path)

            download_name = _download_name(
                original_filename, output_filename, lambda base, ext: f"{base}_images.zip"
            )

        data = read_bytes(output_path)
        return data, download_name
    except Exception as e:
        logger.error(f"DOCX to Image conversion failed: {e}")
        raise e
    finally:
        cleanup_paths(file_path, pdf_path, output_path)


def pptx_to_pdf_task(file_path: str, original_filename: str = None) -> tuple[bytes, str]:
    """Convert a PPTX to PDF and return the bytes."""
    logger.info(f"Starting PPTX to PDF task for {file_path}")

    file_path = resolve_path(file_path)
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Input file not found: {file_path}")

    tmp_dir = ensure_tmp_dir()

    output_filename = f"converted_{uuid.uuid4()}.pdf"
    output_path = os.path.join(tmp_dir, output_filename)

    try:
        generated_pdf = _convert_docx_or_pptx_to_pdf(file_path, tmp_dir)
        shutil.move(generated_pdf, output_path)

        data = read_bytes(output_path)
        download_name = _download_name(
            original_filename, output_filename, lambda base, ext: f"{base}.pdf"
        )
        return data, download_name
    except Exception as e:
        logger.error(f"PPTX to PDF conversion failed: {e}")
        raise e
    finally:
        cleanup_paths(file_path, output_path)


def pptx_to_images_task(
    file_path: str, img_format: str, dpi: int, original_filename: str = None
) -> tuple[bytes, str]:
    """Convert PPTX to PDF first, rasterize slides, return single image or ZIP bytes."""
    logger.info(f"Starting PPTX to images task for {file_path}")

    file_path = resolve_path(file_path)
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Input file not found: {file_path}")

    tmp_dir = ensure_tmp_dir()

    pdf_path = None
    output_filename = ""
    output_path = ""

    try:
        pdf_path = _convert_docx_or_pptx_to_pdf(file_path, tmp_dir)

        images = convert_from_path(pdf_path, dpi=dpi)

        if len(images) == 1:
            output_filename = f"slide_{uuid.uuid4()}.{img_format.lower()}"
            output_path = os.path.join(tmp_dir, output_filename)
            images[0].save(output_path, format=img_format.upper())

            download_name = _download_name(
                original_filename,
                output_filename,
                lambda base, ext: f"{base}.{img_format.lower()}",
            )
        else:
            output_filename = f"slides_{uuid.uuid4()}.zip"
            output_path = os.path.join(tmp_dir, output_filename)
            with zipfile.ZipFile(output_path, "w") as zipf:
                for idx, img in enumerate(images):
                    img_name = f"slide_{idx + 1}.{img_format.lower()}"
                    temp_img_path = os.path.join(tmp_dir, img_name)
                    img.save(temp_img_path, format=img_format.upper())
                    zipf.write(temp_img_path, img_name)
                    os.remove(temp_img_path)

            download_name = _download_name(
                original_filename, output_filename, lambda base, ext: f"{base}_slides.zip"
            )

        data = read_bytes(output_path)
        return data, download_name
    except Exception as e:
        logger.error(f"PPTX to images conversion failed: {e}")
        raise e
    finally:
        cleanup_paths(file_path, pdf_path, output_path)


def merge_pptx_task(file_paths: list[str], original_filename: str = None) -> tuple[bytes, str]:
    """Combine PowerPoint files into one deck and return the bytes."""
    logger.info(f"Starting PPTX merge task for files: {file_paths}")

    from pptx import Presentation

    import copy

    resolved_paths = resolve_paths(file_paths)

    output_filename = f"merged_{uuid.uuid4()}.pptx"
    tmp_dir = ensure_tmp_dir()
    output_path = os.path.join(tmp_dir, output_filename)

    try:
        valid_paths = [fp for fp in resolved_paths if os.path.exists(fp)]
        if not valid_paths:
            raise Exception("No valid PowerPoint files found to merge.")

        prs1 = Presentation(valid_paths[0])

        for fp in valid_paths[1:]:
            prs2 = Presentation(fp)

            for slide in prs2.slides:
                blank_layout = prs1.slide_layouts[6]
                new_slide = prs1.slides.add_slide(blank_layout)

                for shape in slide.shapes:
                    shape_el = copy.deepcopy(shape.element)
                    new_slide.shapes._spTree.append(shape_el)

        prs1.save(output_path)

        data = read_bytes(output_path)
        download_name = _download_name(
            original_filename, output_filename, lambda base, ext: f"{base}_merged{ext}"
        )
        return data, download_name
    except Exception as e:
        logger.error(f"PPTX merge task failed: {e}")
        raise e
    finally:
        for fp in resolved_paths:
            if fp and os.path.exists(fp):
                try:
                    os.remove(fp)
                except Exception:
                    pass
        if os.path.exists(output_path):
            try:
                os.remove(output_path)
            except Exception:
                pass
