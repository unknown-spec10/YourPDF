"""Synchronous endpoint tests: status codes, content types, roundtrips."""
import io
import zipfile

from pypdf import PdfReader, PdfWriter
from PIL import Image


def _make_pdf_bytes(pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_text_pdf_bytes(text: str = "hello world") -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(200, 200))
    c.drawString(50, 100, text)
    c.save()
    return buf.getvalue()


def _make_image_bytes(fmt: str = "JPEG", size=(100, 80)) -> bytes:
    img = Image.new("RGB", size, color=(200, 30, 30))
    buf = io.BytesIO()
    img.save(buf, fmt)
    return buf.getvalue()


def _make_docx_bytes(text: str = "hello") -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pptx_bytes(text: str = "hello") -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tx_box = slide.shapes.add_textbox(left=0, top=0, width=1000000, height=500000)
    tx_box.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf_file(name: str = "doc.pdf", pages: int = 1):
    return {"file": (name, _make_pdf_bytes(pages), "application/pdf")}


def test_health_check(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok", "service": "yourpdf-backend"}


def test_file_size_exceeded(client):
    large_data = b"0" * (31 * 1024 * 1024)
    files = {"file": ("large.pdf", large_data, "application/pdf")}
    response = client.post("/api/compress", files=files, data={"quality": "low"})
    assert response.status_code == 400
    assert "exceeds the maximum limit" in response.json()["detail"]
    assert "30MB" in response.json()["detail"]


def test_compress_invalid_extension(client):
    files = {"file": ("test.txt", b"plain text data", "text/plain")}
    response = client.post("/api/compress", files=files, data={"quality": "low"})
    assert response.status_code == 400
    assert response.json()["detail"] == "Only PDF files are supported."


def test_merge_pdfs_roundtrip(client):
    files = [
        ("files", ("doc1.pdf", _make_pdf_bytes(1), "application/pdf")),
        ("files", ("doc2.pdf", _make_pdf_bytes(2), "application/pdf")),
    ]
    response = client.post("/api/merge", files=files)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "doc1_merged" in response.headers.get("content-disposition", "")
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 3


def test_merge_pdfs_single_file_rejected(client):
    files = [("files", ("doc1.pdf", _make_pdf_bytes(1), "application/pdf"))]
    response = client.post("/api/merge", files=files)
    assert response.status_code == 400
    assert response.json()["detail"] == "At least 2 PDF files are required to merge."


def test_split_custom_roundtrip(client):
    response = client.post(
        "/api/split", files=_pdf_file(pages=3), data={"mode": "custom", "pages_spec": "1-2"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "doc_split" in response.headers.get("content-disposition", "")
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 2


def test_split_missing_spec_rejected(client):
    response = client.post(
        "/api/split", files=_pdf_file(), data={"mode": "custom", "pages_spec": ""}
    )
    assert response.status_code == 400


def test_split_all_multi_page_returns_zip(client):
    response = client.post(
        "/api/split", files=_pdf_file(pages=2), data={"mode": "all", "pages_spec": ""}
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert response.headers.get("content-disposition", "").endswith('.zip"')
    names = zipfile.ZipFile(io.BytesIO(response.content)).namelist()
    assert names == ["page_1.pdf", "page_2.pdf"]


def test_images_to_pdf_roundtrip(client):
    files = [
        ("files", ("img1.png", _make_image_bytes("PNG"), "image/png")),
        ("files", ("img2.jpg", _make_image_bytes("JPEG"), "image/jpeg")),
    ]
    response = client.post("/api/images-to-pdf", files=files)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "compiled" in response.headers.get("content-disposition", "")
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 2


def test_images_to_pdf_rejects_non_image(client):
    files = [("files", ("doc.pdf", _make_pdf_bytes(), "application/pdf"))]
    response = client.post("/api/images-to-pdf", files=files)
    assert response.status_code == 400


def test_extract_text_inline_json(client):
    response = client.post(
        "/api/extract-text",
        files={"file": ("doc.pdf", _make_text_pdf_bytes("hello world"), "application/pdf")},
        data={"export_txt": "false"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/json"
    assert "hello world" in response.json()["text"]


def test_extract_text_export_txt(client):
    response = client.post(
        "/api/extract-text",
        files={"file": ("doc.pdf", _make_text_pdf_bytes("hello world"), "application/pdf")},
        data={"export_txt": "true"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/plain")
    assert "doc_extracted" in response.headers.get("content-disposition", "")
    assert b"hello world" in response.content


def test_protect_pdf_success(client):
    pdf_bytes = _make_pdf_bytes()
    files = {"file": ("doc.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/protect", files=files, data={"password": "secret"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "doc_protected" in response.headers.get("content-disposition", "")
    reader = PdfReader(io.BytesIO(response.content))
    assert reader.is_encrypted


def test_unlock_pdf_success(client):
    from app.pdf_security import protect_pdf_bytes

    protected = protect_pdf_bytes(_make_pdf_bytes(), "secret")
    files = {"file": ("doc_protected.pdf", protected, "application/pdf")}
    response = client.post("/api/unlock", files=files, data={"password": "secret"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    reader = PdfReader(io.BytesIO(response.content))
    assert len(reader.pages) == 1


def test_unlock_pdf_wrong_password(client):
    from app.pdf_security import protect_pdf_bytes

    protected = protect_pdf_bytes(_make_pdf_bytes(), "secret")
    files = {"file": ("doc_protected.pdf", protected, "application/pdf")}
    response = client.post("/api/unlock", files=files, data={"password": "wrong"})
    assert response.status_code == 400


def test_rotate_pdf_roundtrip(client):
    response = client.post("/api/rotate", files=_pdf_file(), data={"rotation": "90"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "doc_rotated" in response.headers.get("content-disposition", "")
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 1


def test_rotate_pdf_invalid_angle(client):
    response = client.post("/api/rotate", files=_pdf_file(), data={"rotation": "0"})
    assert response.status_code == 400


def test_organize_pdf_roundtrip(client):
    response = client.post("/api/organize", files=_pdf_file(pages=2), data={"page_order": "2,1"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "doc_organized" in response.headers.get("content-disposition", "")
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 2


def test_organize_pdf_missing_order(client):
    # Empty required Form fields are rejected by the framework with 422
    # (unchanged pre-existing behavior).
    response = client.post("/api/organize", files=_pdf_file(), data={"page_order": ""})
    assert response.status_code == 422


def test_watermark_pdf_roundtrip(client):
    response = client.post(
        "/api/watermark",
        files=_pdf_file(),
        data={"text": "COPY", "color": "red", "opacity": "0.5", "rotation": "45"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "doc_watermarked" in response.headers.get("content-disposition", "")
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 1


def test_add_page_numbers_roundtrip(client):
    response = client.post(
        "/api/add-page-numbers", files=_pdf_file(), data={"position": "bottom-right"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "doc_numbered" in response.headers.get("content-disposition", "")
    assert len(PdfReader(io.BytesIO(response.content)).pages) == 1


def test_image_invalid_extension(client):
    files = {"file": ("test.pdf", b"%PDF-1.4 dummy", "application/pdf")}
    response = client.post("/api/image/compress", files=files, data={"quality": "75"})
    assert response.status_code == 400
    assert response.json()["detail"] == "Only JPG, JPEG, PNG, and WEBP image files are supported."


def _image_file(name: str = "photo.jpg"):
    return {"file": (name, _make_image_bytes("JPEG"), "image/jpeg")}


def test_compress_image_roundtrip(client):
    response = client.post("/api/image/compress", files=_image_file(), data={"quality": "80"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/jpeg"
    assert "photo_compressed" in response.headers.get("content-disposition", "")
    Image.open(io.BytesIO(response.content)).verify()


def test_compress_image_bad_quality(client):
    response = client.post("/api/image/compress", files=_image_file(), data={"quality": "150"})
    assert response.status_code == 400


def test_resize_image_roundtrip(client):
    response = client.post(
        "/api/image/resize",
        files=_image_file(),
        data={"width": "50", "height": "40", "maintain_aspect": "false"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/jpeg"
    img = Image.open(io.BytesIO(response.content))
    assert img.size == (50, 40)


def test_resize_image_missing_dims(client):
    response = client.post("/api/image/resize", files=_image_file(), data={})
    assert response.status_code == 400


def test_crop_image_roundtrip(client):
    response = client.post(
        "/api/image/crop",
        files=_image_file(),
        data={"x": "10", "y": "10", "width": "40", "height": "30"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/jpeg"
    assert Image.open(io.BytesIO(response.content)).size == (40, 30)


def test_crop_image_bad_dims(client):
    response = client.post(
        "/api/image/crop",
        files=_image_file(),
        data={"x": "0", "y": "0", "width": "0", "height": "10"},
    )
    assert response.status_code == 400


def test_convert_image_roundtrip(client):
    response = client.post(
        "/api/image/convert", files=_image_file(), data={"target_format": "png"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/png"
    assert response.headers.get("content-disposition", "").endswith('.png"')


def test_convert_image_bad_format(client):
    response = client.post(
        "/api/image/convert", files=_image_file(), data={"target_format": "svg"}
    )
    assert response.status_code == 400


def test_rotate_image_roundtrip(client):
    response = client.post("/api/image/rotate", files=_image_file(), data={"angle": "90"})
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/jpeg"
    assert "photo_rotated" in response.headers.get("content-disposition", "")


def test_rotate_image_bad_angle(client):
    response = client.post("/api/image/rotate", files=_image_file(), data={"angle": "45"})
    assert response.status_code == 400


def test_watermark_image_roundtrip(client):
    response = client.post(
        "/api/image/watermark",
        files=_image_file(),
        data={"text": "CONFIDENTIAL", "color": "red", "opacity": "0.5", "position": "center"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/jpeg"
    Image.open(io.BytesIO(response.content)).verify()


def test_watermark_image_empty_text(client):
    # Empty required Form fields are rejected by the framework with 422
    # (unchanged pre-existing behavior).
    response = client.post(
        "/api/image/watermark", files=_image_file(), data={"text": ""}
    )
    assert response.status_code == 422


def test_merge_docx_roundtrip(client):
    files = [
        ("files", ("doc1.docx", _make_docx_bytes("one"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("files", ("doc2.docx", _make_docx_bytes("two"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ]
    response = client.post("/api/office/merge-docx", files=files)
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "doc1_merged" in response.headers.get("content-disposition", "")


def test_merge_docx_single_file_rejected(client):
    files = [
        ("files", ("doc1.docx", _make_docx_bytes("one"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ]
    response = client.post("/api/office/merge-docx", files=files)
    assert response.status_code == 400


def test_merge_pptx_roundtrip(client):
    files = [
        ("files", ("slides1.pptx", _make_pptx_bytes("one"), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
        ("files", ("slides2.pptx", _make_pptx_bytes("two"), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
    ]
    response = client.post("/api/office/merge-pptx", files=files)
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert "slides1_merged" in response.headers.get("content-disposition", "")


def test_merge_pptx_single_file_rejected(client):
    files = [
        ("files", ("slides1.pptx", _make_pptx_bytes("one"), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
    ]
    response = client.post("/api/office/merge-pptx", files=files)
    assert response.status_code == 400


def test_pdf_info_invalid_extension(client):
    files = {"file": ("test.txt", b"plain text data", "text/plain")}
    response = client.post("/api/pdf-info", files=files)
    assert response.status_code == 400
    assert response.json()["detail"] == "Only PDF files are supported."


def test_images_to_pdf_task_direct():
    """Direct unit check of the sync images-to-pdf function (no HTTP)."""
    import os
    import tempfile

    from app.tasks import images_to_pdf_task

    with tempfile.TemporaryDirectory() as tmp:
        p1 = os.path.join(tmp, "a.png")
        p2 = os.path.join(tmp, "b.jpg")
        with open(p1, "wb") as f:
            f.write(_make_image_bytes("PNG"))
        with open(p2, "wb") as f:
            f.write(_make_image_bytes("JPEG"))
        data, name = images_to_pdf_task([p1, p2], "original.png")
        assert data.startswith(b"%PDF")
        assert name == "original_compiled.pdf"
        assert not os.path.exists(p1) and not os.path.exists(p2)
